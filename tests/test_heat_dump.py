"""Metriky mařeného tepla v porovnání scénářů, detailu profilu a plánu.

Mařené (zahozené) teplo je teplo z KGJ, které soustava neodebere. Vidět má
být, kolik ho je celkem, v kolika hodinách, kolik průměrně v hodině maření
a jak dlouhé jsou souvislé úseky maření (průměr a nejdelší).
"""

import io

import openpyxl
import pandas as pd
import pytest

from app_extract import load_app
from conftest import SOLVER_TIME_LIMIT, make_df, make_params, make_uses
from opt_core import (calculate_smoothness_metrics, heat_dump_stats,
                      run_optimization_with_profile)


def stats(values):
    return heat_dump_stats(pd.DataFrame({'Zahozené teplo [MW]': values}))


def test_runs_of_heat_dumping():
    """Úseky 2, 1 a 3 h: celkem 1,5 MWh v 6 h, souvisle průměrně 2 h."""
    assert stats([0, 0.5, 0.5, 0, 0.2, 0, 0.1, 0.1, 0.1, 0]) == pytest.approx(
        {'total_mwh': 1.5, 'hours': 6, 'avg_mwh': 0.25,
         'avg_run_hours': 2.0, 'max_run_hours': 3})


def test_dumping_at_both_ends_of_horizon():
    assert stats([0.3, 0.3, 0, 0, 0.3]) == pytest.approx(
        {'total_mwh': 0.9, 'hours': 3, 'avg_mwh': 0.3,
         'avg_run_hours': 1.5, 'max_run_hours': 2})


def test_solver_noise_is_not_dumping():
    assert stats([1e-9, 0.0, 2e-7])['hours'] == 0


@pytest.mark.parametrize('frame', [
    pd.DataFrame({'Zahozené teplo [MW]': [0.0] * 5}),
    pd.DataFrame({'KGJ on': [1.0] * 5}),               # sloupec vubec neni
])
def test_no_dumping_gives_zeros(frame):
    assert heat_dump_stats(frame) == {'total_mwh': 0.0, 'hours': 0,
                                      'avg_mwh': 0.0, 'avg_run_hours': 0.0,
                                      'max_run_hours': 0}


# ── exporty ─────────────────────────────────────────────────────────

@pytest.fixture(scope='module')
def dumping():
    """KGJ jede naplno v drahých hodinách (0–5 a 10–19), soustava bere 0,1 MW.

    Výkon 0,975 MW, min. zatížení 100 %: v každé hodině provozu se zahodí
    zhruba 0,875 MW, úseky maření mají 6 a 10 h.
    """
    prices = [300.0] * 6 + [0.0] * 4 + [300.0] * 10 + [0.0] * 4
    df = make_df(hours=24, ee_price=prices, heat_demand=0.1)
    r = run_optimization_with_profile(df, make_params(), make_uses(),
                                      profile_type='free',
                                      time_limit=SOLVER_TIME_LIMIT)
    assert r is not None
    res = r['res']
    s = heat_dump_stats(res)
    assert (s['hours'], s['avg_run_hours'], s['max_run_hours']) == (16, 8.0, 10)
    assert s['avg_mwh'] == pytest.approx(0.975 - 0.1, abs=0.005)
    assert s['total_mwh'] == pytest.approx(16 * s['avg_mwh'])
    scenarios = {'free': {'result': r, 'smoothness': calculate_smoothness_metrics(res),
                          'profile_name': 'FREE'}}
    return scenarios, s


COMPARISON = {
    'Mařené teplo [MWh]':      lambda s: f"{s['total_mwh']:.1f}",
    'Hodin maření':            lambda s: s['hours'],
    'Ø maření v hodině [MWh]': lambda s: f"{s['avg_mwh']:.2f}",
    'Ø souvislé maření [h]':   lambda s: f"{s['avg_run_hours']:.1f}",
    'Max souvislé maření [h]': lambda s: s['max_run_hours'],
}


def test_comparison_table_has_heat_dump(dumping):
    scenarios, s = dumping
    row = load_app().create_scenario_comparison_df(scenarios).iloc[0]
    for col, value in COMPARISON.items():
        assert row[col] == value(s), col


def test_scenarios_excel_has_heat_dump(dumping):
    scenarios, s = dumping
    data = load_app().to_excel_scenarios(scenarios)
    ws = openpyxl.load_workbook(io.BytesIO(data))['Porovnání scénářů']
    header = [c.value for c in ws[1]]
    for col, value in COMPARISON.items():
        assert ws.cell(row=2, column=header.index(col) + 1).value == value(s), col


def test_operating_plan_overview_has_heat_dump(dumping):
    scenarios, s = dumping
    data = load_app().to_excel_operating_plan(scenarios['free'], 'free')
    ws = openpyxl.load_workbook(io.BytesIO(data))['Přehled']
    kpis = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value
            for r in range(1, ws.max_row + 1)}
    assert kpis['Mařené teplo celkem [MWh]'] == pytest.approx(s['total_mwh'])
    assert kpis['Hodiny maření tepla [h]'] == s['hours']
    assert kpis['Průměr v hodině maření [MWh]'] == pytest.approx(s['avg_mwh'])
    assert kpis['Průměrné souvislé maření [h]'] == pytest.approx(s['avg_run_hours'])
    assert kpis['Nejdelší souvislé maření [h]'] == s['max_run_hours']
