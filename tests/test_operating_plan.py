"""Testy provozního plánu — měsíční mřížka provozu a export do Excelu.

Layout se řídí předlohou `plan_pr.xlsx`, kterou uživatel dodal: dny ve
sloupcích, hodiny v řádcích, P = provoz, X = klid.
"""

import ast
from pathlib import Path
import io
import types

import numpy as np
import openpyxl
import pandas as pd
import pytest
from xlsxwriter.utility import xl_col_to_name

import opt_core
from conftest import SOLVER_TIME_LIMIT, make_params, make_uses, with_ramp
from opt_core import (HOUR_LABELS, MONTH_NAMES_FULL, build_month_grid,
                      run_optimization_with_profile)

REPO = opt_core.__file__.rsplit('/', 1)[0]

# Bunka A5 je hodina 0, sloupec B je den 1.
ROW_HOUR0, COL_DAY1 = 5, 2


def _export_module():
    """Vytáhne exportní funkce z app.py.

    app.py je Streamlit skript, takže ho nejde importovat — načteme z něj
    přes AST jen ty funkce, které testujeme.
    """
    src = Path(f'{REPO}/app.py').read_text(encoding='utf-8')
    mod = types.ModuleType('app_export')
    mod.__dict__.update({
        'io': io, 'pd': pd, 're': __import__('re'),
        'xl_col_to_name': xl_col_to_name,
        'HOUR_LABELS': HOUR_LABELS, 'MONTH_NAMES_FULL': MONTH_NAMES_FULL,
        'build_month_grid': build_month_grid,
    })
    wanted = {'_wb_formats', '_safe_sheet', '_write_sheet', '_round_numeric',
              'build_parameters_df', '_write_month_sheet',
              'to_excel_operating_plan'}
    for node in ast.parse(src).body:
        if isinstance(node, ast.FunctionDef) and node.name in wanted:
            exec(compile(ast.Module([node], []), '<app>', 'exec'), mod.__dict__)
    missing = wanted - set(mod.__dict__)
    assert not missing, f'v app.py chybí {missing}'
    return mod


@pytest.fixture(scope='module')
def solved():
    """Měsíc PEAK provozu se zapnutou rampou — vzniknou doběhové hodiny."""
    T = 24 * 31
    df = pd.DataFrame({
        'datetime': pd.date_range('2026-01-01', periods=T, freq='h'),
        'ee_price': 90 + 40 * np.sin(np.arange(T) * 2 * np.pi / 24),
        'gas_price': np.full(T, 38.0),
        'Poptávka po teple (MW)': np.full(T, 0.5),
        'FVE (MW)': np.zeros(T),
    })
    p = with_ramp(make_params(k_th=0.605, k_min=0.5, k_min_runtime=4,
                              k_start_cost=150.0, h_cover=0.9,
                              shortfall_penalty=500.0), 12.7, 8.9)
    r = run_optimization_with_profile(df, p, make_uses(), profile_type='peak',
                                      time_limit=SOLVER_TIME_LIMIT)
    assert r is not None
    return r['res'], p


@pytest.fixture(scope='module')
def workbook(solved):
    res, p = solved
    data = _export_module().to_excel_operating_plan(
        {'result': {'res': res}}, 'peak', params=p, uses=make_uses())
    return openpyxl.load_workbook(io.BytesIO(data))


# ── Jádro zadání ─────────────────────────────────────────────────────

def test_tail_hour_is_marked_off(solved, workbook):
    """Doběhová hodina musí být X, i když v ní jednotka ještě dodává teplo.

    Zadání znělo, že sjezdová hodina už v téhle tabulce figurovat nemá.
    """
    res, _ = solved
    ws = workbook['LEDEN']
    times = pd.to_datetime(res['Čas'])

    tails = res.index[(res['KGJ stop'] > 0.5) & (res['KGJ doběh [MW_th]'] > 1e-9)]
    assert len(tails) > 0, 'test má smysl jen když nějaký doběh nastal'

    for i in tails:
        ts = times.iloc[i]
        cell = ws.cell(row=ROW_HOUR0 + ts.hour, column=COL_DAY1 + ts.day - 1)
        assert res['KGJ [MW_th]'][i] > 0, 'doběh má dodávat teplo'
        assert cell.value == 'X', f'{ts} je doběh, má být X, je {cell.value!r}'


def test_grid_reads_commitment_not_output(solved):
    """Mřížka odpovídá sloupci KGJ on, hodinu po hodině."""
    res, _ = solved
    days, grid = build_month_grid(res, 1)
    times = pd.to_datetime(res['Čas'])

    for i in range(len(res)):
        ts = times.iloc[i]
        expected = 'P' if res['KGJ on'][i] > 0.5 else 'X'
        assert grid[ts.hour][ts.day - 1] == expected, f'neshoda v {ts}'


# ── Popisky hodin ────────────────────────────────────────────────────

def test_hour_labels_are_time_ranges(workbook):
    ws = workbook['LEDEN']
    labels = [ws.cell(row=5 + h, column=1).value for h in range(24)]
    assert labels[0] == '00:00-01:00'
    assert labels[-1] == '23:00-24:00'
    assert labels == HOUR_LABELS
    # navazujici intervaly, zadna mezera
    for a, b in zip(labels, labels[1:]):
        assert a.split('-')[1] == b.split('-')[0]


def test_first_row_is_midnight_hour(solved, workbook):
    """Řádek 00:00-01:00 nese data z hodiny 0, ne z hodiny 1.

    Samotné popisky můžou vypadat správně, i když je mřížka posunutá.
    """
    res, _ = solved
    ws = workbook['LEDEN']
    times = pd.to_datetime(res['Čas'])

    for day in (1, 15, 31):
        for hour in (0, 12, 23):
            match = res.index[(times.dt.day == day) & (times.dt.hour == hour)]
            expected = 'P' if res['KGJ on'][match[0]] > 0.5 else 'X'
            cell = ws.cell(row=ROW_HOUR0 + hour, column=COL_DAY1 + day - 1)
            assert cell.value == expected, f'den {day}, hodina {hour}'


# ── Layout podle předlohy ────────────────────────────────────────────

def test_layout_matches_template(workbook):
    ws = workbook['LEDEN']
    assert ws['A1'].value == 'zpět na úvod'
    assert ws['A2'].value == 'měsíc'
    assert ws['B2'].value == 'LEDEN'
    assert ws['A3'].value == 'den'
    assert ws['A4'].value == 'hodina'
    assert ws['B3'].value == 1
    assert ws['B4'].value == 'provoz'
    assert ws.cell(row=3, column=1 + 31).value == 31, 'leden má 31 dní'


def test_counts_are_live_formulas(workbook):
    """Počty musí být vzorce, aby ruční přepsání buňky přepočítalo součty."""
    ws = workbook['LEDEN']
    assert ws['A30'].value == 'počet P'
    assert ws['A31'].value == 'počet X'
    assert ws['B30'].value == '=COUNTIF(B5:B28,"P")'
    assert ws['B31'].value == '=COUNTIF(B5:B28,"X")'
    assert ws['B34'].value == '=SUM(B30:AF30)'
    assert ws['B35'].value == '=SUM(B31:AF31)'
    assert 'Provoz' in ws['C34'].value
    assert 'klidu' in ws['C35'].value


def test_conditional_formatting_present(workbook):
    """Barvy jdou z podmíněného formátování, jako v předloze."""
    ws = workbook['LEDEN']
    found = {}
    for rng in ws.conditional_formatting:
        for rule in rng.rules:
            for val in ('P', 'X', 'F'):
                if rule.formula and f'"{val}"' in rule.formula[0]:
                    found[val] = rule.dxf.fill.end_color.rgb[-6:].upper()
    assert found.get('P') == 'C6EFCE', 'provoz zeleně'
    assert found.get('X') == 'FFC7CE', 'klid červeně'
    assert found.get('F') == 'FFEB9C', 'žluté pravidlo zůstává pro ruční použití'


def test_sheet_order_and_names(workbook):
    names = workbook.sheetnames
    assert names[0] == 'Přehled', 'přehled musí být první list'
    assert names[1] == 'PEAK', 'pak hodinový rozpad profilu'
    assert 'LEDEN' in names


def test_month_sheets_link_back(workbook):
    ws = workbook['LEDEN']
    assert ws['A1'].hyperlink is not None
    assert 'Přehled' in str(ws['A1'].hyperlink.location or ws['A1'].hyperlink.target)


def test_overview_links_to_months(workbook):
    ws = workbook['Přehled']
    links = [c.value for row in ws.iter_rows() for c in row
             if c.hyperlink is not None]
    assert 'LEDEN' in links


# ── Okrajové případy ─────────────────────────────────────────────────

def _grid_from(times, on_flags):
    res = pd.DataFrame({'Čas': times, 'KGJ on': on_flags})
    return build_month_grid(res, pd.to_datetime(times[0]).month)


def test_dst_duplicate_hour_collapses():
    """Zdvojená hodina na konci října dá jednu buňku — P, když běžela aspoň jednou."""
    times = ['2026-10-25 02:00', '2026-10-25 02:00', '2026-10-25 03:00']
    days, grid = _grid_from(times, [0, 1, 0])
    assert grid[2][24] == 'P', 'běžela v jedné z dvojice, má být P'

    days, grid = _grid_from(times, [0, 0, 0])
    assert grid[2][24] == 'X'


def test_partial_month_leaves_blanks():
    """Dny mimo analyzované období zůstanou prázdné, ne X."""
    times = pd.date_range('2026-01-01', periods=48, freq='h')
    days, grid = build_month_grid(
        pd.DataFrame({'Čas': times, 'KGJ on': np.zeros(48)}), 1)

    assert len(days) == 31, 'mřížka pokrývá celý kalendářní měsíc'
    assert grid[0][0] == 'X', 'první den data má'
    assert grid[0][5] is None, 'šestý den v datech není -> prázdno, ne X'


def test_grid_is_empty_for_month_without_data():
    times = pd.date_range('2026-01-01', periods=24, freq='h')
    days, grid = build_month_grid(
        pd.DataFrame({'Čas': times, 'KGJ on': np.ones(24)}), 7)
    assert days == [] and grid == []
