"""Testy exportu porovnání scénářů — přehled hodin pod hlavní tabulkou.

Zadání: na prvním listu pod výsledky každého profilu má být tabulka,
v jakých hodinách ten profil jede.
"""

import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import SOLVER_TIME_LIMIT, make_params, make_uses
from opt_core import (HOUR_LABELS, build_hour_month_matrix,
                      calculate_smoothness_metrics,
                      run_optimization_with_profile)

PROFILES = ('p', 'peak')


def _export_module():
    from app_extract import load_app
    return load_app()


@pytest.fixture(scope='module')
def scenarios():
    T = 24 * 31
    t = pd.date_range('2026-01-01', periods=T, freq='h')
    df = pd.DataFrame({
        'datetime': t,
        'ee_price': 90 + 40 * np.sin(np.arange(T) * 2 * np.pi / 24),
        'gas_price': np.full(T, 38.0),
        'Poptávka po teple (MW)': np.full(T, 1.0),
        'FVE (MW)': np.zeros(T),
    })
    p = make_params(k_th=0.605, k_min=0.5, k_min_runtime=4, k_start_cost=50.0)
    out = {}
    for prof in PROFILES:
        r = run_optimization_with_profile(df, p, make_uses(),
                                          profile_type=prof,
                                          time_limit=SOLVER_TIME_LIMIT)
        out[prof] = {'result': r, 'profile': prof,
                     'smoothness': calculate_smoothness_metrics(r['res'])}
    return out, p


@pytest.fixture(scope='module')
def wb(scenarios):
    scen, p = scenarios
    data = _export_module().to_excel_scenarios(scen, params=p,
                                               uses=make_uses())
    return openpyxl.load_workbook(io.BytesIO(data)), scen


def _rows_with(ws, text):
    return [r for r in range(1, ws.max_row + 1)
            if isinstance(ws.cell(row=r, column=1).value, str)
            and text in ws.cell(row=r, column=1).value]


def test_first_sheet_is_the_comparison(wb):
    book, _ = wb
    assert book.sheetnames[0] == 'Porovnání scénářů'


def test_every_profile_has_its_hour_table(wb):
    book, scen = wb
    ws = book['Porovnání scénářů']
    for prof in scen:
        assert _rows_with(ws, f'Profil {prof.upper()}'), \
            f'chybi prehled hodin pro {prof}'


def test_hour_tables_sit_below_the_comparison(wb):
    """Tabulky patri pod hlavni tabulku, ne nad ni."""
    book, scen = wb
    ws = book['Porovnání scénářů']
    head = _rows_with(ws, 'V jakých hodinách profily jedou')
    assert head, 'chybi nadpis sekce'
    first = min(min(_rows_with(ws, f'Profil {p.upper()}')) for p in scen)
    assert head[0] < first
    assert head[0] > 1 + len(scen), 'sekce zacina uz v hlavni tabulce'


def test_hour_tables_do_not_overlap(wb):
    """Kazdy profil ma vlastni blok 24 radku + hlavicka + soucet."""
    book, scen = wb
    ws = book['Porovnání scénářů']
    heads = sorted(min(_rows_with(ws, f'Profil {p.upper()}')) for p in scen)
    for a, b in zip(heads, heads[1:]):
        assert b - a >= 27, f'bloky se prekryvaji: {a} -> {b}'


def test_values_match_the_computed_matrix(wb):
    book, scen = wb
    ws = book['Porovnání scénářů']
    for prof, s in scen.items():
        hdr = min(_rows_with(ws, f'Profil {prof.upper()}')) + 1
        exp = build_hour_month_matrix(s['result']['res'])
        assert ws.cell(row=hdr, column=2).value == 'I'
        for h in range(24):
            assert ws.cell(row=hdr + 1 + h, column=1).value == HOUR_LABELS[h]
            assert ws.cell(row=hdr + 1 + h, column=2).value == \
                int(exp.iloc[h, 0]), (prof, h)


def test_counts_are_numbers_not_text(wb):
    """Sloupce maji z hlavni tabulky textovy format — cisla ho musi prebit."""
    book, scen = wb
    ws = book['Porovnání scénářů']
    prof = next(iter(scen))
    hdr = min(_rows_with(ws, f'Profil {prof.upper()}')) + 1
    vals = [ws.cell(row=hdr + 1 + h, column=2).value for h in range(24)]
    assert all(isinstance(v, int) for v in vals), vals
    assert ws.cell(row=hdr + 1, column=2).number_format == '0'


def test_profile_p_window_visible_in_export(wb):
    """Leden u profilu P je 07-23, takze krajni hodiny musi byt nuly."""
    book, scen = wb
    ws = book['Porovnání scénářů']
    hdr = min(_rows_with(ws, 'Profil P')) + 1
    jan = [ws.cell(row=hdr + 1 + h, column=2).value for h in range(24)]
    assert sum(jan[:7]) == 0, 'pred 07:00 ma byt klid'
    assert jan[23] == 0, 'po 23:00 ma byt klid'
    assert sum(jan[7:23]) > 0, 'uvnitr okna ma neco bezet'
