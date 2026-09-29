"""Změna času v provozním plánu.

Data jsou v místním čase s letním časem, stejně jako dodaná FWD křivka:
29. 3. 2026 chybí hodina 02:00 a 25. 10. 2026 je 02:00 dvakrát. V měsíční
mřížce z toho bylo prázdné bílé políčko (jaro), resp. změna nebyla vidět
vůbec (podzim). Obě políčka teď mají vlastní značení.
"""

import io

import openpyxl
import pandas as pd
import pytest

from opt_core import HOUR_LABELS, build_month_grid, find_dst_hours


def local_year(year=2026):
    """Hodinová osa v místním čase CET/CEST jako naivní razítka."""
    t = pd.date_range(f'{year}-01-01', f'{year}-12-31 23:00', freq='h',
                      tz='Europe/Prague')
    return pd.Series(t.tz_localize(None))


def res_from(times, on=1.0):
    return pd.DataFrame({'Čas': times, 'KGJ on': on})


@pytest.fixture(scope='module')
def res():
    return res_from(local_year())


# ── detekce ──────────────────────────────────────────────────────────

def test_axis_looks_like_the_fwd_file(res):
    """Pojistka: testovací osa má stejný tvar jako dodaná data."""
    t = res['Čas']
    assert len(t) == 8760
    assert pd.Timestamp('2026-03-29 02:00') not in set(t)
    assert (t == pd.Timestamp('2026-10-25 02:00')).sum() == 2


def test_spring_gap_found(res):
    assert find_dst_hours(res, 3) == {(29, 2): 'gap'}


def test_autumn_double_found(res):
    assert find_dst_hours(res, 10) == {(25, 2): 'double'}


@pytest.mark.parametrize('month', [1, 2, 4, 5, 6, 7, 8, 9, 11, 12])
def test_other_months_have_nothing(res, month):
    assert find_dst_hours(res, month) == {}


def test_other_years_use_their_own_sunday():
    """Pravidlo je poslední neděle v měsíci, ne pevné datum."""
    r = res_from(local_year(2027))
    assert find_dst_hours(r, 3) == {(28, 2): 'gap'}
    assert find_dst_hours(r, 10) == {(31, 2): 'double'}


def test_naive_24h_axis_marks_nothing():
    """Data bez letního času (každý den 24 h) žádnou změnu nemají."""
    t = pd.Series(pd.date_range('2026-01-01', periods=8760, freq='h'))
    r = res_from(t)
    assert find_dst_hours(r, 3) == {}
    assert find_dst_hours(r, 10) == {}


def test_ordinary_data_hole_is_not_dst():
    """Díra v datech v jiný den se za změnu času vydávat nesmí."""
    t = pd.Series(pd.date_range('2026-03-01', '2026-03-31 23:00', freq='h'))
    t = t[t != pd.Timestamp('2026-03-10 02:00')]
    assert find_dst_hours(res_from(t), 3) == {}


def test_period_starting_after_the_gap_is_not_dst():
    """Období začínající 29. 3. v 05:00 nemá 02:00 kvůli začátku, ne kvůli času."""
    t = pd.Series(pd.date_range('2026-03-29 05:00', '2026-03-31 23:00',
                                freq='h'))
    assert find_dst_hours(res_from(t), 3) == {}


def test_grid_still_leaves_gap_empty(res):
    """Mřížka samotná se nemění — značení přidává až export."""
    _days, grid = build_month_grid(res, 3)
    assert grid[2][28] is None


# ── export ───────────────────────────────────────────────────────────

@pytest.fixture(scope='module')
def plan(res):
    from app_extract import load_app
    app = load_app()
    t = res['Čas']
    r = res.assign(**{
        'Hodinový zisk [€]': 0.0, 'Dodáno tepla [MW]': 0.0,
        'EE z KGJ [MW]': 0.0, 'EE export [MW]': 0.0, 'Shortfall [MW]': 0.0,
        # provoz jen 01:00-04:00, at je videt P i X kolem zmeny casu
        'KGJ on': [1.0 if 1 <= ts.hour < 4 else 0.0 for ts in t],
    })
    data = app.to_excel_operating_plan({'result': {'res': r}}, 'prom26')
    return openpyxl.load_workbook(io.BytesIO(data)), app


def _cell(ws, day, hour):
    return ws.cell(row=5 + hour, column=1 + day)


def test_spring_cell_is_not_white(plan):
    wb, app = plan
    c = _cell(wb['BŘEZEN'], 29, 2)
    assert c.value == app.DST_GAP_MARK
    assert c.fill.fgColor.rgb.endswith(app.DST_COLOR.lstrip('#').upper()), \
        'jarni zmena casu ma mit vlastni vypln'


def test_spring_mark_is_not_counted_as_p_or_x(plan):
    """'ZČ' nesmí spadnout do počtu P ani X — ta hodina neexistuje."""
    wb, app = plan
    mark = app.DST_GAP_MARK
    assert 'P' not in mark and 'X' not in mark and 'F' not in mark


def test_spring_cell_has_comment(plan):
    wb, _ = plan
    c = _cell(wb['BŘEZEN'], 29, 2)
    assert c.comment is not None and 'letní čas' in c.comment.text


def test_autumn_cell_keeps_value_and_gets_border(plan):
    """Podzim: P/X zůstává, změnu času nese rámeček a poznámka."""
    wb, app = plan
    c = _cell(wb['ŘÍJEN'], 25, 2)
    assert c.value == 'P'
    assert c.border.left.style not in (None, 'thin'), 'chybi zvyrazneny ramecek'
    assert c.comment is not None and 'zimní čas' in c.comment.text


def test_ordinary_cells_unchanged(plan):
    wb, _ = plan
    ws = wb['BŘEZEN']
    assert _cell(ws, 28, 2).value == 'P'
    assert _cell(ws, 28, 2).comment is None
    assert _cell(ws, 29, 3).value == 'P'
    assert _cell(ws, 29, 5).value == 'X'


def test_legend_only_where_needed(plan):
    wb, _ = plan

    def legend(ws):
        return [ws.cell(row=r, column=3).value for r in range(37, 40)
                if ws.cell(row=r, column=3).value]

    assert any('letní čas' in v for v in legend(wb['BŘEZEN']))
    assert any('zimní čas' in v for v in legend(wb['ŘÍJEN']))
    assert legend(wb['LEDEN']) == []


def test_hour_labels_untouched(plan):
    wb, _ = plan
    ws = wb['BŘEZEN']
    assert [ws.cell(row=5 + h, column=1).value for h in range(24)] == HOUR_LABELS
