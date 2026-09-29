"""Testy přehledu 'v kterých hodinách profil jede'.

Tabulka je 24 × měsíce, hodnota = počet dní, kdy KGJ v té hodině běžela.
Stejný výpočet jde do porovnání scénářů i do provozního plánu, proto je
v jádře a ne v `app.py`.
"""

import io

import numpy as np
import openpyxl
import pandas as pd
import pytest

from conftest import SOLVER_TIME_LIMIT, make_params, make_uses
from opt_core import (HOUR_LABELS, build_hour_month_matrix, build_month_grid,
                      run_optimization_with_profile)


def _res(hours, start='2026-01-01 00:00', on=None):
    """Minimalni ramec ve tvaru, jaky tabulka ocekava."""
    t = pd.date_range(start, periods=hours, freq='h')
    return pd.DataFrame({'Čas': t,
                         'KGJ on': [1.0] * hours if on is None else on})


# ── výpočet ──────────────────────────────────────────────────────────

def test_shape_and_labels():
    m = build_hour_month_matrix(_res(24 * 31))
    assert list(m.index) == HOUR_LABELS
    assert list(m.columns) == ['I']
    assert m.shape == (24, 1)


def test_counts_are_days_not_hours():
    """Leden v provozu 24/7 -> v kazde hodine 31 dni."""
    m = build_hour_month_matrix(_res(24 * 31))
    assert (m['I'] == 31).all()
    assert m['I'].sum() == 24 * 31


def test_zero_where_profile_never_runs():
    """Hodiny mimo okno musi vyjit na nulu, ne chybet."""
    t = pd.date_range('2026-01-01', periods=24 * 31, freq='h')
    on = [(1.0 if 8 <= ts.hour < 20 else 0.0) for ts in t]
    m = build_hour_month_matrix(pd.DataFrame({'Čas': t, 'KGJ on': on}))
    assert (m['I'][8:20] == 31).all()
    assert m['I'][0:8].sum() == 0 and m['I'][20:24].sum() == 0
    assert m['I'].sum() == 12 * 31


def test_several_months_keep_roman_order():
    m = build_hour_month_matrix(_res(24 * 120))     # I-IV
    assert list(m.columns) == ['I', 'II', 'III', 'IV']


def test_partial_month_counts_only_its_days():
    m = build_hour_month_matrix(_res(24 * 10))      # 10 dni ledna
    assert (m['I'] == 10).all()


def test_dst_duplicate_hour_counts_as_one_day():
    """Pri prechodu na zimni cas ma 25. 10. dve razitka 02:00.

    Bez slouceni by ten den v hodine 02 vysel jako dva dny a sloupec by
    prerostl pocet dni v mesici.
    """
    t = list(pd.date_range('2026-10-25 00:00', periods=3, freq='h'))
    t.insert(3, pd.Timestamp('2026-10-25 02:00'))   # duplicitni hodina
    d = pd.DataFrame({'Čas': t, 'KGJ on': [1.0] * 4})
    m = build_hour_month_matrix(d)
    assert m.loc[HOUR_LABELS[2], 'X'] == 1, 'duplicitni hodina se ma slucovat'
    assert m['X'].max() <= 31


def test_tail_hour_is_not_counted_as_running():
    """Cte se 'KGJ on', tedy nasazeni — stejne jako mesicni mrizka."""
    t = pd.date_range('2026-01-01', periods=4, freq='h')
    d = pd.DataFrame({'Čas': t, 'KGJ on': [1.0, 1.0, 0.0, 0.0],
                      'KGJ doběh [MW_th]': [0.0, 0.0, 0.3, 0.0]})
    m = build_hour_month_matrix(d)
    assert m['I'].sum() == 2, 'dobeh se nesmi pocitat jako provoz'


def test_matches_month_grid():
    """Souctem po hodinach musi tabulka sednout na mesicni mrizku."""
    t = pd.date_range('2026-03-01', periods=24 * 31, freq='h')
    rng = np.random.default_rng(7)
    on = rng.integers(0, 2, len(t)).astype(float)
    d = pd.DataFrame({'Čas': t, 'KGJ on': on})
    m = build_hour_month_matrix(d)
    _days, grid = build_month_grid(d, 3)
    for h in range(24):
        assert m.iloc[h, 0] == sum(1 for v in grid[h] if v == 'P'), h


# ── zápis do sešitu ──────────────────────────────────────────────────

def _export_module():
    from app_extract import load_app
    return load_app()


@pytest.fixture(scope='module')
def plan_wb():
    """Provozni plan pro profil P nad dvema mesici realistickych dat."""
    T = 24 * 59                                   # leden + unor
    t = pd.date_range('2026-01-01', periods=T, freq='h')
    df = pd.DataFrame({
        'datetime': t,
        'ee_price': 90 + 40 * np.sin(np.arange(T) * 2 * np.pi / 24),
        'gas_price': np.full(T, 38.0),
        'Poptávka po teple (MW)': np.full(T, 1.0),
        'FVE (MW)': np.zeros(T),
    })
    p = make_params(k_th=0.605, k_min=0.5, k_min_runtime=4, k_start_cost=50.0)
    r = run_optimization_with_profile(df, p, make_uses(), profile_type='p',
                                      time_limit=SOLVER_TIME_LIMIT)
    data = _export_module().to_excel_operating_plan(
        {'result': {'res': r['res']}}, 'p', params=p, uses=make_uses())
    return openpyxl.load_workbook(io.BytesIO(data), data_only=False), r['res']


def _find_row(ws, text):
    for row in range(1, ws.max_row + 1):
        v = ws.cell(row=row, column=1).value
        if isinstance(v, str) and text in v:
            return row
    return None


def test_plan_overview_has_hour_table(plan_wb):
    wb, _ = plan_wb
    ws = wb['Přehled']
    assert _find_row(ws, 'V jakých hodinách profil jede') is not None, \
        'v provoznim planu chybi prehled hodin'


def test_plan_hour_table_matches_data(plan_wb):
    wb, res = plan_wb
    ws = wb['Přehled']
    head = _find_row(ws, 'Profil P')
    assert head is not None
    hdr = head + 1
    assert ws.cell(row=hdr, column=1).value == 'hodina'
    months = [ws.cell(row=hdr, column=c).value for c in (2, 3)]
    assert months == ['I', 'II']

    expected = build_hour_month_matrix(res)
    for h in range(24):
        assert ws.cell(row=hdr + 1 + h, column=1).value == HOUR_LABELS[h]
        for j in range(2):
            assert ws.cell(row=hdr + 1 + h, column=2 + j).value == \
                int(expected.iloc[h, j]), (h, j)


def test_plan_hour_table_totals_are_live_formulas(plan_wb):
    """Soucty jsou vzorce, at rucni oprava bunky prepocita radek."""
    wb, _ = plan_wb
    ws = wb['Přehled']
    hdr = _find_row(ws, 'Profil P') + 1
    total = ws.cell(row=hdr + 25, column=1).value
    assert total == 'celkem [h]'
    f = ws.cell(row=hdr + 25, column=2).value
    assert isinstance(f, str) and f.startswith('=SUM('), f


def test_plan_hour_table_sits_below_month_links(plan_wb):
    """Tabulka patri pod mesicni prehled, ne pred nej."""
    wb, _ = plan_wb
    ws = wb['Přehled']
    assert _find_row(ws, 'Měsíc') < _find_row(ws, 'V jakých hodinách profil jede')


def test_profile_p_window_in_plan(plan_wb):
    """Leden 07-23, unor 06-22 — mimo okno musi byt nuly."""
    wb, _ = plan_wb
    ws = wb['Přehled']
    hdr = _find_row(ws, 'Profil P') + 1
    jan = [ws.cell(row=hdr + 1 + h, column=2).value for h in range(24)]
    feb = [ws.cell(row=hdr + 1 + h, column=3).value for h in range(24)]
    assert sum(jan[:7]) == 0 and sum(jan[23:]) == 0, 'leden mimo 07-23'
    assert sum(feb[:6]) == 0 and sum(feb[22:]) == 0, 'unor mimo 06-22'
